"""Versioned, section-based report storage, validation, revisions, and export."""

import io
import json
import re
import shutil
import threading
import time
import uuid
from functools import wraps
from pathlib import Path
from typing import Any

from rcp.config import data_dir
from rcp.objects import (
    ClaimBundle,
    EvidenceLink,
    ExperimentPlan,
    ExperimentResultSet,
    Hypothesis, PaperCard, ResearchSynthesis,
    ReportDocument,
    ReportFigure,
    ReportMetadata,
    ReportReference,
    ReportRevision,
    ReportSection,
    ResultBundle,
    utc_now,
)


class ReportConflict(ValueError):
    pass


_LOCKS: dict[str, threading.RLock] = {}
_LOCKS_GUARD = threading.Lock()


def _serialized(function):
    """Make optimistic version validation and file replacement atomic per run."""
    @wraps(function)
    def wrapper(run_id: str, *args, **kwargs):
        with _LOCKS_GUARD:
            lock = _LOCKS.setdefault(run_id, threading.RLock())
        with lock:
            return function(run_id, *args, **kwargs)
    return wrapper


def _run_dir(run_id: str) -> Path:
    return data_dir() / "runs" / run_id


def _document_path(run_id: str) -> Path:
    return _run_dir(run_id) / "report_document.json"


def _revision_dir(run_id: str) -> Path:
    return _run_dir(run_id) / "report_revisions"


def _section_id(title: str, used: set[str]) -> str:
    base = re.sub(r"[^a-z0-9]+", "-", title.lower()).strip("-") or "section"
    candidate = base
    index = 2
    while candidate in used:
        candidate = f"{base}-{index}"
        index += 1
    used.add(candidate)
    return candidate


def _safe_figure_path(run_id: str, value: str) -> Path | None:
    if not value:
        return None
    root = _run_dir(run_id).resolve()
    candidate = Path(value)
    candidate = candidate.resolve() if candidate.is_absolute() else (root / candidate).resolve()
    if not candidate.is_relative_to(root) or not candidate.is_file():
        return None
    return candidate


def _comparison_heatmap(
    run_id: str, experiment_results: ExperimentResultSet | None
) -> ReportFigure | None:
    comparisons = experiment_results.comparisons if experiment_results else []
    if not comparisons:
        return None
    from PIL import Image as PILImage, ImageDraw, ImageFont

    pairs = list(dict.fromkeys(item.pair_id for item in comparisons))
    metrics = list(dict.fromkeys(item.metric for item in comparisons))
    cell_width, cell_height = 118, 42
    label_width, header_height = 245, 68
    width = label_width + cell_width * len(pairs) + 20
    height = header_height + cell_height * len(metrics) + 58
    image = PILImage.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    try:
        regular = ImageFont.truetype("DejaVuSans.ttf", 13)
        small = ImageFont.truetype("DejaVuSans.ttf", 11)
        bold = ImageFont.truetype("DejaVuSans-Bold.ttf", 13)
    except OSError:
        regular = small = bold = ImageFont.load_default()
    draw.text((12, 10), "Matched baseline → candidate change", fill="#17202a", font=bold)
    for column, pair in enumerate(pairs):
        x = label_width + column * cell_width
        draw.text((x + 5, 38), pair.removeprefix("pair-"), fill="#4b5563", font=small)
    lookup = {(item.metric, item.pair_id): item for item in comparisons}
    for row, metric in enumerate(metrics):
        y = header_height + row * cell_height
        draw.text((12, y + 13), metric, fill="#374151", font=regular)
        for column, pair in enumerate(pairs):
            x = label_width + column * cell_width
            item = lookup.get((metric, pair))
            if item is None:
                fill, label = "#f3f4f6", "—"
            else:
                fill = "#d1fae5" if item.candidate_better is True else (
                    "#fee2e2" if item.candidate_better is False else "#e5e7eb"
                )
                label = "—" if item.percent_delta is None else f"{item.percent_delta:+.2f}%"
            draw.rectangle((x, y, x + cell_width - 5, y + cell_height - 5), fill=fill, outline="#ffffff")
            draw.text((x + 9, y + 12), label, fill="#1f2937", font=regular)
    draw.text(
        (12, height - 35),
        "Green: candidate better under preregistered direction · Red: worse · Gray: neutral/unscored",
        fill="#6b7280", font=small,
    )
    relative = Path("figures") / "comparison_heatmap.png"
    destination = _run_dir(run_id) / relative
    destination.parent.mkdir(parents=True, exist_ok=True)
    image.save(destination, format="PNG", optimize=True)
    return ReportFigure(
        id="matched-comparison-heatmap",
        title="Matched comparison heatmap",
        path=str(relative),
        caption=(
            "Baseline-to-candidate percentage changes for every matched case and primary metric. "
            "Color follows the preregistered metric direction."
        ),
        evidence_ids=[item.id for item in comparisons],
    )


def render_markdown(document: ReportDocument) -> str:
    lines = [f"# {document.metadata.title}", ""]
    if document.metadata.authors:
        lines += ["**Authors:** " + ", ".join(document.metadata.authors), ""]
    for section in document.sections:
        lines += [f"## {section.title}", "", section.content.rstrip(), ""]
    if document.figures:
        lines += ["## Figures", ""]
        for figure in document.figures:
            lines += [f"![{figure.caption or figure.title}]({figure.path})", ""]
    return "\n".join(lines).rstrip() + "\n"


def _parse_legacy(run_id: str, text: str) -> ReportDocument:
    title_match = re.search(r"^#\s+(.+)$", text, re.MULTILINE)
    title = title_match.group(1).strip() if title_match else f"Research report {run_id}"
    matches = list(re.finditer(r"^##\s+(.+)$", text, re.MULTILINE))
    sections: list[ReportSection] = []
    used: set[str] = set()
    if matches:
        for index, match in enumerate(matches):
            start = match.end()
            end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
            section_title = match.group(1).strip()
            content = text[start:end].strip()
            sid = _section_id(section_title, used)
            sections.append(
                ReportSection(
                    id=sid, title=section_title, kind=sid, content=content,
                    generated=False, generated_content=content,
                )
            )
    else:
        content = text.strip()
        sections.append(ReportSection(id="report", title="Report", content=content, generated=False, generated_content=content))
    return ReportDocument(
        run_id=run_id,
        metadata=ReportMetadata(
            title=title, run_id=run_id,
            disclosures=["Legacy Markdown report migrated to the structured report workspace."],
        ),
        sections=sections,
        migrated_from_legacy=True,
    )


def load_report_document(run_id: str, persist_migration: bool = False) -> ReportDocument:
    path = _document_path(run_id)
    if path.exists():
        return ReportDocument.model_validate_json(path.read_text())
    legacy = _run_dir(run_id) / "report.md"
    if not legacy.exists():
        raise FileNotFoundError(f"report for run {run_id} has not been drafted")
    document = _parse_legacy(run_id, legacy.read_text())
    if persist_migration:
        backup = _run_dir(run_id) / "report.legacy.md"
        if not backup.exists():
            shutil.copyfile(legacy, backup)
        _save_new(document, "Legacy report migrated")
    return document


def _write_document(document: ReportDocument) -> None:
    root = _run_dir(document.run_id)
    root.mkdir(parents=True, exist_ok=True)
    document_path = _document_path(document.run_id)
    document_tmp = root / f".{document_path.name}.{uuid.uuid4().hex}.tmp"
    markdown_path = root / "report.md"
    markdown_tmp = root / f".{markdown_path.name}.{uuid.uuid4().hex}.tmp"
    document_tmp.write_text(document.model_dump_json(indent=2))
    markdown_tmp.write_text(render_markdown(document))
    document_tmp.replace(document_path)
    markdown_tmp.replace(markdown_path)


def _archive(document: ReportDocument, summary: str) -> ReportRevision:
    root = _revision_dir(document.run_id)
    root.mkdir(parents=True, exist_ok=True)
    existing = sorted(root.glob(f"v{document.version:06d}-*.json"))
    if existing:
        try:
            return ReportRevision.model_validate(json.loads(existing[-1].read_text())["revision"])
        except (KeyError, ValueError, json.JSONDecodeError):
            pass
    revision = ReportRevision(
        id=f"v{document.version:06d}-{int(time.time() * 1000)}-{uuid.uuid4().hex[:4]}",
        version=document.version,
        summary=summary,
    )
    payload = {"revision": revision.model_dump(), "document": document.model_dump()}
    (root / f"{revision.id}.json").write_text(json.dumps(payload, indent=2))
    return revision


def _save_new(document: ReportDocument, summary: str) -> ReportDocument:
    document.updated_at = utc_now()
    _write_document(document)
    _archive(document, summary)
    return document


def seed_report_document(
    run_id: str,
    title: str,
    bodies: dict[str, str],
    cards: list[PaperCard],
    claim_bundle: ClaimBundle,
    result_bundle: ResultBundle | None,
    experiment_plan: ExperimentPlan | None = None,
    experiment_results: ExperimentResultSet | None = None,
    research_synthesis: ResearchSynthesis | None = None,
    selected_hypothesis: Hypothesis | None = None,
) -> ReportDocument:
    references = [
        ReportReference(
            id=f"ref-{index}", paper_id=card.id,
            citation=f"{card.title} ({card.year or 'n.d.'})",
            doi=card.doi, url=card.url,
        )
        for index, card in enumerate(cards[:12], 1)
    ]
    reference_by_paper = {reference.paper_id: reference.id for reference in references}
    metric_links = [
        EvidenceLink(
            id=f"metric:{name}", kind="metric", label=f"{name} = {value}",
            source=(result_bundle.metric_methods.get(name, "result_bundle") if result_bundle else "result_bundle"),
        )
        for name, value in (result_bundle.metrics if result_bundle else {}).items()
    ]
    claim_links = [
        EvidenceLink(id=f"claim:{index}", kind="claim", label=claim.statement, source=claim.evidence)
        for index, claim in enumerate(claim_bundle.claims, 1)
    ]
    comparison_links = [
        EvidenceLink(
            id=comparison.id, kind="comparison",
            label=(
                f"{comparison.metric}: {comparison.baseline_value:g} → "
                f"{comparison.candidate_value:g} (Δ {comparison.absolute_delta:+g})"
            ),
            source=f"{comparison.baseline_case_id},{comparison.candidate_case_id}",
        )
        for comparison in (experiment_results.comparisons if experiment_results else [])
    ]
    case_links = [
        EvidenceLink(
            id=f"case:{case.id}", kind="experiment-case", label=case.label or case.id,
            source=case.model_name,
        )
        for case in (experiment_plan.cases if experiment_plan else [])
    ]
    plan_links = [
        EvidenceLink(
            id=f"plan:{experiment_plan.id}", kind="experiment-plan",
            label=experiment_plan.description or experiment_plan.id,
            source=experiment_plan.hypothesis_id,
        )
    ] if experiment_plan else []
    protocol_links = [
        EvidenceLink(
            id=f"protocol:{experiment_plan.analysis_protocol.id}", kind="analysis-protocol",
            label=(
                f"{experiment_plan.analysis_protocol.id} · thermal threshold "
                f"{experiment_plan.analysis_protocol.thermal_threshold_degC:g} °C · HVAC screen "
                f"{experiment_plan.analysis_protocol.hvac_power_screen_min_W:g}–"
                f"{experiment_plan.analysis_protocol.hvac_power_screen_max_W:g} W"
            ),
            source=(experiment_results.analysis_protocol_sha256 if experiment_results else ""),
        )
    ] if experiment_plan else []
    quality_links = [
        EvidenceLink(
            id=f"quality:{check.id}", kind="study-quality-check",
            label=f"{check.status}: {check.message}", source=check.observed,
        )
        for check in (
            experiment_results.quality_report.checks
            if experiment_results and experiment_results.quality_report else []
        )
    ]
    raw_series_links = [
        EvidenceLink(
            id=f"raw-series:{result.case_id or result.spec_id}", kind="raw-series",
            label=(
                f"{result.case_id or result.spec_id} raw CSV · "
                f"{result.result_file_integrity} · sha256 "
                f"{(result.result_file_sha256 or 'unavailable')[:16]}"
            ),
            source=result.result_file or "",
        )
        for result in (experiment_results.cases if experiment_results else [])
        if result.status == "ok"
    ]
    finding_links = [
        EvidenceLink(
            id=finding.id, kind="literature-finding", label=finding.statement,
            source=finding.paper_id,
        )
        for card in cards for finding in card.findings
    ]
    gap_links = [
        EvidenceLink(
            id=gap.id, kind="research-gap", label=gap.statement,
            source=",".join(gap.paper_ids),
        )
        for gap in (research_synthesis.gaps if research_synthesis else [])
    ]
    conflict_links = [
        EvidenceLink(
            id=conflict.id, kind=f"research-{conflict.kind}", label=conflict.statement,
            source=",".join(conflict.finding_ids),
        )
        for conflict in (research_synthesis.conflicts if research_synthesis else [])
    ]
    hypothesis_links = [
        EvidenceLink(
            id=f"hypothesis:{selected_hypothesis.id}", kind="hypothesis",
            label=selected_hypothesis.statement,
            source=",".join([
                *selected_hypothesis.supporting_gap_ids,
                *selected_hypothesis.supporting_conflict_ids,
                *selected_hypothesis.supporting_paper_ids,
            ]),
        )
    ] if selected_hypothesis else []
    citation_ids = [ref.id for ref in references]
    metric_ids = [link.id for link in metric_links]
    claim_evidence_ids = list(dict.fromkeys(
        evidence_id for claim in claim_bundle.claims
        for evidence_id in [*claim.evidence_ids, *claim.comparison_ids, *[f"case:{case_id}" for case_id in claim.case_ids]]
    ))
    claims_text = "\n".join(
        f"- **{claim.statement}** — evidence: {claim.evidence} (confidence: {claim.confidence})"
        for claim in claim_bundle.claims
    ) or "No supported claims were generated."
    refs_text = "\n".join(
        f"{index}. {ref.citation} {ref.url or (('https://doi.org/' + ref.doi) if ref.doi else '')}".rstrip()
        for index, ref in enumerate(references, 1)
    ) or "No references were available."
    synthesis_citation_ids = list(dict.fromkeys(
        reference_by_paper[paper_id]
        for item in [
            *(research_synthesis.gaps if research_synthesis else []),
            *(research_synthesis.conflicts if research_synthesis else []),
        ]
        for paper_id in item.paper_ids
        if paper_id in reference_by_paper
    ))
    synthesis_evidence_ids = [
        *[link.id for link in gap_links], *[link.id for link in conflict_links],
        *[link.id for link in hypothesis_links],
    ]
    synthesis_text = "No structured research synthesis was available for this run."
    if research_synthesis:
        gap_text = "\n".join(
            f"- **{gap.category.title()} gap ({gap.id}):** {gap.statement}"
            for gap in research_synthesis.gaps
        ) or "- No supported gaps were identified from the available abstracts."
        conflict_text = "\n".join(
            f"- **{conflict.kind.title()} ({conflict.id}):** {conflict.statement}"
            for conflict in research_synthesis.conflicts
        ) or "- No cross-paper contradictions or tensions passed validation."
        hypothesis_text = (
            f"\n\n**Selected hypothesis (rank {selected_hypothesis.rank}, "
            f"score {selected_hypothesis.rank_score:.1f}/100):** {selected_hypothesis.statement}"
            if selected_hypothesis else ""
        )
        synthesis_text = (
            f"Evidence coverage: {research_synthesis.papers_with_findings} of "
            f"{research_synthesis.papers_total} papers yielded abstract-grounded findings "
            f"({research_synthesis.finding_count} findings).\n\n"
            f"### Research gaps\n{gap_text}\n\n### Cross-paper conflicts\n{conflict_text}"
            f"{hypothesis_text}"
        )
    disclosures = list(dict.fromkeys([
        *(result_bundle.warnings if result_bundle else []),
        *(warning for result in (experiment_results.cases if experiment_results else []) for warning in result.warnings),
        *(experiment_results.warnings if experiment_results else []),
    ])) or [
        "Simulation results and model assumptions must be interpreted within their stated validation status."
    ]
    disclosure = " ".join(disclosures)
    section_specs = [
        ("introduction", "Introduction", bodies.get("introduction", ""), citation_ids, []),
        ("research-synthesis", "Research Synthesis", synthesis_text, synthesis_citation_ids, synthesis_evidence_ids),
        ("method", "Method", bodies.get("method", "") + f"\n\n**Simulation disclosure:** {disclosure}", [], [link.id for link in [*plan_links, *protocol_links, *case_links, *raw_series_links]]),
        ("results-discussion", "Results and Discussion", bodies.get("results_discussion", ""), [], [*metric_ids, *[link.id for link in comparison_links], *[link.id for link in quality_links], *[link.id for link in raw_series_links]]),
        ("conclusion", "Conclusion", bodies.get("conclusion", ""), [], [*metric_ids, *[link.id for link in comparison_links]]),
        ("claims-evidence", "Claims and Evidence", claims_text, [], [*claim_evidence_ids, *[link.id for link in claim_links]]),
        ("references", "References", refs_text, citation_ids, []),
    ]
    sections = [
        ReportSection(
            id=sid, title=heading, kind=sid, content=content, generated_content=content,
            citation_ids=refs, evidence_ids=evidence,
        )
        for sid, heading, content, refs, evidence in section_specs
    ]
    document = ReportDocument(
        run_id=run_id,
        metadata=ReportMetadata(
            title=title, run_id=run_id,
            disclosures=disclosures,
        ),
        sections=sections, references=references,
        figures=[figure] if (figure := _comparison_heatmap(run_id, experiment_results)) else [],
        evidence_links=[
            *metric_links, *claim_links, *comparison_links, *case_links, *plan_links,
            *protocol_links, *quality_links, *raw_series_links, *finding_links,
            *gap_links, *conflict_links, *hypothesis_links,
        ],
    )
    return _save_new(document, "Initial generated report")


def _editable(run_id: str, expected_version: int | None) -> ReportDocument:
    document = load_report_document(run_id, persist_migration=True)
    if expected_version is not None and document.version != expected_version:
        raise ReportConflict(f"report version conflict: expected {expected_version}, current {document.version}")
    return document


def _commit(document: ReportDocument, summary: str) -> ReportDocument:
    _archive(document, f"Before: {summary}")
    document.version += 1
    document.updated_at = utc_now()
    _write_document(document)
    return document


@_serialized
def update_section(
    run_id: str, section_id: str, changes: dict[str, Any], expected_version: int | None = None
) -> ReportDocument:
    document = _editable(run_id, expected_version)
    section = next((item for item in document.sections if item.id == section_id), None)
    if section is None:
        raise KeyError(section_id)
    if section.locked and not (set(changes) == {"locked"} and changes.get("locked") is False):
        raise ReportConflict("locked sections must be unlocked before editing")
    for field in ("title", "content", "locked", "citation_ids", "evidence_ids"):
        if field in changes:
            setattr(section, field, changes[field])
    if "content" in changes:
        section.generated = False
        section.stale = False
    section.updated_at = utc_now()
    return _commit(document, f"Updated section {section.title}")


@_serialized
def add_section(
    run_id: str, title: str, content: str = "", after_id: str | None = None,
    expected_version: int | None = None,
) -> ReportDocument:
    document = _editable(run_id, expected_version)
    sid = _section_id(title, {section.id for section in document.sections})
    section = ReportSection(id=sid, title=title, content=content, generated=False, generated_content=content)
    if after_id:
        index = next((i for i, item in enumerate(document.sections) if item.id == after_id), len(document.sections) - 1)
        document.sections.insert(index + 1, section)
    else:
        document.sections.append(section)
    return _commit(document, f"Added section {title}")


@_serialized
def delete_section(
    run_id: str, section_id: str, expected_version: int | None = None, force: bool = False
) -> ReportDocument:
    document = _editable(run_id, expected_version)
    section = next((item for item in document.sections if item.id == section_id), None)
    if section is None:
        raise KeyError(section_id)
    if section.locked and not force:
        raise ReportConflict("locked sections cannot be deleted without force=true")
    document.sections = [item for item in document.sections if item.id != section_id]
    return _commit(document, f"Deleted section {section.title}")


@_serialized
def reorder_sections(run_id: str, section_ids: list[str], expected_version: int | None = None) -> ReportDocument:
    document = _editable(run_id, expected_version)
    current = {section.id: section for section in document.sections}
    if set(section_ids) != set(current) or len(section_ids) != len(current):
        raise ValueError("section_ids must contain every section exactly once")
    document.sections = [current[sid] for sid in section_ids]
    return _commit(document, "Reordered sections")


@_serialized
def regenerate_section(
    run_id: str, section_id: str, expected_version: int | None = None, force: bool = False
) -> ReportDocument:
    document = _editable(run_id, expected_version)
    section = next((item for item in document.sections if item.id == section_id), None)
    if section is None:
        raise KeyError(section_id)
    if section.locked:
        raise ReportConflict("locked sections cannot be regenerated")
    if section.content != section.generated_content and not force:
        raise ReportConflict("section has manual edits; pass force=true to overwrite them")
    section.content = section.generated_content
    section.generated = True
    section.stale = False
    section.updated_at = utc_now()
    return _commit(document, f"Regenerated section {section.title}")


@_serialized
def update_metadata(run_id: str, changes: dict[str, Any], expected_version: int | None = None) -> ReportDocument:
    document = _editable(run_id, expected_version)
    for field in ("title", "authors", "abstract", "keywords", "disclosures"):
        if field in changes:
            setattr(document.metadata, field, changes[field])
    return _commit(document, "Updated report metadata")


def list_revisions(run_id: str) -> list[ReportRevision]:
    root = _revision_dir(run_id)
    revisions: list[ReportRevision] = []
    if root.exists():
        for path in root.glob("*.json"):
            try:
                revisions.append(ReportRevision.model_validate(json.loads(path.read_text())["revision"]))
            except (KeyError, ValueError, json.JSONDecodeError):
                continue
    return sorted(revisions, key=lambda revision: (revision.created_at, revision.id), reverse=True)


@_serialized
def restore_revision(run_id: str, revision_id: str, expected_version: int | None = None) -> ReportDocument:
    current = _editable(run_id, expected_version)
    path = _revision_dir(run_id) / f"{revision_id}.json"
    if not path.exists():
        raise KeyError(revision_id)
    restored = ReportDocument.model_validate(json.loads(path.read_text())["document"])
    if restored.run_id != run_id:
        raise ValueError("revision belongs to a different run")
    _archive(current, f"Before restoring {revision_id}")
    restored.version = current.version + 1
    restored.updated_at = utc_now()
    _write_document(restored)
    return restored


def validate_report(document: ReportDocument) -> dict[str, Any]:
    issues: list[dict[str, str | None]] = []
    if not document.metadata.title.strip():
        issues.append({"severity": "error", "code": "missing-title", "message": "Report title is required.", "section_id": None})
    if not document.metadata.disclosures:
        issues.append({"severity": "error", "code": "missing-simulation-disclosure", "message": "Add a simulation/model credibility disclosure before export.", "section_id": None})
    if not document.metadata.authors:
        issues.append({"severity": "warning", "code": "incomplete-authors", "message": "Report authors have not been entered.", "section_id": None})
    if not document.metadata.abstract.strip():
        issues.append({"severity": "warning", "code": "incomplete-abstract", "message": "Report abstract is incomplete.", "section_id": None})
    reference_ids = {ref.id for ref in document.references}
    evidence_ids = {link.id for link in document.evidence_links}
    for section in document.sections:
        if not section.title.strip():
            issues.append({"severity": "error", "code": "missing-section-title", "message": "Every report section requires a title.", "section_id": section.id})
        if not section.content.strip():
            issues.append({"severity": "warning", "code": "empty-section", "message": "Section content is empty.", "section_id": section.id})
        missing_refs = sorted(set(section.citation_ids) - reference_ids)
        if missing_refs:
            issues.append({"severity": "error", "code": "missing-citation", "message": f"Unknown citation IDs: {', '.join(missing_refs)}", "section_id": section.id})
        missing_evidence = sorted(set(section.evidence_ids) - evidence_ids)
        if missing_evidence:
            issues.append({"severity": "error", "code": "missing-evidence", "message": f"Unknown evidence IDs: {', '.join(missing_evidence)}", "section_id": section.id})
        claim_text = section.content.strip()
        if (
            section.kind == "claims-evidence" and claim_text
            and claim_text != "No supported claims were generated."
            and not section.evidence_ids
        ):
            issues.append({"severity": "error", "code": "unsupported-claims", "message": "Claims section has no structured evidence links.", "section_id": section.id})
        if section.stale:
            issues.append({"severity": "warning", "code": "stale-generated-section", "message": "Generated content is stale relative to its evidence.", "section_id": section.id})
        for citation in re.findall(r"\[(\d+)\]", section.content):
            if not 1 <= int(citation) <= len(document.references):
                issues.append({"severity": "error", "code": "missing-citation", "message": f"Inline citation [{citation}] has no matching reference.", "section_id": section.id})
    for figure in document.figures:
        missing_evidence = sorted(set(figure.evidence_ids) - evidence_ids)
        if missing_evidence:
            issues.append({"severity": "error", "code": "missing-figure-evidence", "message": f"Figure {figure.id} has unknown evidence IDs: {', '.join(missing_evidence)}", "section_id": None})
        if not _safe_figure_path(document.run_id, figure.path):
            issues.append({"severity": "error", "code": "missing-figure-file", "message": f"Figure {figure.id} has no readable file inside the run directory.", "section_id": None})
    errors = sum(issue["severity"] == "error" for issue in issues)
    warnings = sum(issue["severity"] == "warning" for issue in issues)
    return {"valid": errors == 0, "errors": errors, "warnings": warnings, "issues": issues, "version": document.version}


def export_report(document: ReportDocument, format_name: str) -> tuple[bytes, str, str]:
    validation = validate_report(document)
    if not validation["valid"]:
        raise ValueError("report validation failed: " + "; ".join(issue["message"] for issue in validation["issues"] if issue["severity"] == "error"))
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", document.metadata.title).strip("-").lower()[:50] or "report"
    if format_name == "md":
        return render_markdown(document).encode(), "text/markdown; charset=utf-8", f"{slug}.md"
    if format_name == "docx":
        from docx import Document
        from docx.shared import Inches

        output = io.BytesIO()
        word = Document()
        word.core_properties.title = document.metadata.title
        word.add_heading(document.metadata.title, level=0)
        if document.metadata.authors:
            word.add_paragraph(", ".join(document.metadata.authors))
        for section in document.sections:
            word.add_heading(section.title, level=1)
            for line in section.content.splitlines():
                stripped = re.sub(r"\*\*(.*?)\*\*", r"\1", line).strip()
                if stripped.startswith("- "):
                    word.add_paragraph(stripped[2:], style="List Bullet")
                elif re.match(r"^\d+\. ", stripped):
                    word.add_paragraph(re.sub(r"^\d+\. ", "", stripped), style="List Number")
                elif stripped:
                    word.add_paragraph(stripped)
        if document.figures:
            word.add_heading("Figures", level=1)
            for figure in document.figures:
                path = _safe_figure_path(document.run_id, figure.path)
                if path:
                    word.add_picture(str(path), width=Inches(6))
                    word.add_paragraph(figure.caption or figure.title)
        word.save(output)
        return output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", f"{slug}.docx"
    if format_name == "pdf":
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer
        from xml.sax.saxutils import escape

        output = io.BytesIO()
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle("RCPTitle", parent=styles["Title"], alignment=TA_CENTER, spaceAfter=8 * mm)
        pdf = SimpleDocTemplate(output, pagesize=A4, rightMargin=22 * mm, leftMargin=22 * mm, topMargin=20 * mm, bottomMargin=20 * mm)
        story = [Paragraph(escape(document.metadata.title), title_style)]
        if document.metadata.authors:
            story += [Paragraph(escape(", ".join(document.metadata.authors)), styles["Normal"]), Spacer(1, 5 * mm)]
        for section in document.sections:
            story += [Paragraph(escape(section.title), styles["Heading1"]), Spacer(1, 2 * mm)]
            for paragraph in re.split(r"\n\s*\n|\n(?=[-*] )", section.content):
                cleaned = re.sub(r"\*\*(.*?)\*\*", r"\1", paragraph).strip()
                if cleaned:
                    story += [Paragraph(escape(cleaned).replace("\n", "<br/>"), styles["BodyText"]), Spacer(1, 2 * mm)]
        if document.figures:
            story += [Paragraph("Figures", styles["Heading1"])]
            for figure in document.figures:
                path = _safe_figure_path(document.run_id, figure.path)
                if path:
                    image = Image(str(path))
                    image._restrictSize(165 * mm, 110 * mm)
                    story += [image, Paragraph(escape(figure.caption or figure.title), styles["BodyText"]), Spacer(1, 3 * mm)]
        pdf.build(story)
        return output.getvalue(), "application/pdf", f"{slug}.pdf"
    raise ValueError("format must be one of: md, docx, pdf")
